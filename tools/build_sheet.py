"""Write an Excel workbook of every Prelims question on the site (GS I and CSAT).
Run:  python3 tools/build_sheet.py OUT.xlsx      (needs openpyxl)
"""
import json, pathlib, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
def read(name, var):
    t = (ROOT / name).read_text(encoding="utf-8")
    t = t[t.index(f"window.{var}=") + len(f"window.{var}="):].strip().rstrip(";")
    return json.loads(t)

def sheet(wb, title, data, passages):
    ws = wb.create_sheet(title)
    head = ["ID", "Year", "Q no", "Subject", "Topic"] + (["Passage / directions"] if passages else []) + \
           ["Question", "Option (a)", "Option (b)", "Option (c)", "Option (d)", "Official answer", "Answer text", "Answer key"]
    ws.append(head)
    for r in data["qs"]:
        y, n, s, t, q, o, a = r[:7]
        ans = "Dropped" if a == "x" else a
        row = [f"{y}-{n}", y, n, s, t] + ([data["passages"].get(r[7], "") if len(r) > 7 and r[7] else ""] if passages else []) + \
              [q, *o, ans, "" if a == "x" else o["abcd".index(a)], data["years"].get(str(y), {}).get("key", "")]
        ws.append(row)
    widths = {"ID": 10, "Year": 7, "Q no": 6, "Subject": 22, "Topic": 26, "Passage / directions": 60, "Question": 70,
              "Official answer": 9, "Answer text": 30, "Answer key": 22}
    for i, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 28)
        c = ws.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="C8442F")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
    return len(data["qs"])

wb = Workbook(); wb.remove(wb.active)
csat = read("prelims-csat.js", "CSAT_SEARCH"); gs = read("prelims-gs1.js", "PRE_SEARCH")
n1 = sheet(wb, "GS I", gs, False); n2 = sheet(wb, "CSAT", csat, True)
info = wb.create_sheet("About", 0)
for line in ["PYQ Astra: UPSC Prelims questions with official answers",
             f"GS I: {n1} questions, years {min(gs['years'])}-{max(gs['years'])}",
             f"CSAT: {n2} questions, years {', '.join(sorted(csat['years'], reverse=True))}",
             "Generated from prelims-gs1.js and prelims-csat.js by tools/build_sheet.py.",
             "Each tab has filters on the header row: filter by Year, Subject or Topic."]:
    info.append([line])
info.column_dimensions["A"].width = 90; info["A1"].font = Font(bold=True, size=14)
wb.save(sys.argv[1]); print(f"wrote {sys.argv[1]}: GS I {n1}, CSAT {n2}")
